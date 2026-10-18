from typing import List, Optional
import logging
from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.core.mail import EmailMessage
from django.core.validators import validate_email
from .models import Empresa, Equipamento, FotoTriagem, Funcionario, Peca, SolicitacaoSenha, Tecnico, Triagem
from .dtos import EmpresaDTO, EquipamentoDTO, FuncionarioDTO, TecnicoDTO, TriagemDTO, StatusPeca
import os
from django.conf import settings
from docx import Document
from io import BytesIO
from datetime import date
from django.contrib.auth.hashers import make_password, check_password, identify_hasher
from django.utils import timezone
from django.utils.crypto import constant_time_compare
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Count
from PIL import Image, ImageOps, UnidentifiedImageError
import uuid
from django.core.cache import cache
from .partsurfer import buscar_pecas_por_pn

logger = logging.getLogger(__name__)


# ==========================================
# SERVIÇO DE EMPRESA
# ==========================================
class EmpresaService:
    @staticmethod
    def salvar(dto: EmpresaDTO) -> EmpresaDTO:
        if dto.id_empresa:
            empresa = Empresa.objects.get(pk=dto.id_empresa)
            empresa.cnpj = dto.cnpj
            empresa.razao_social = dto.razao_social
            empresa.tipo = dto.tipo
            empresa.save()
        else:
            empresa = Empresa.objects.create(
                cnpj=dto.cnpj,
                razao_social=dto.razao_social,
                tipo=dto.tipo
            )
        
        return EmpresaDTO(
            id_empresa=empresa.id_empresa,
            cnpj=empresa.cnpj,
            razao_social=empresa.razao_social,
            tipo=empresa.tipo
        )

    @staticmethod
    def listar(tipo_filtro: Optional[int] = None) -> List[EmpresaDTO]:
        # Se receber um filtro, busca apenas aquele tipo. Senão, busca todos.
        if tipo_filtro is not None:
            empresas = Empresa.objects.filter(tipo=tipo_filtro)
        else:
            empresas = Empresa.objects.all()
            
        return [
            EmpresaDTO(
                id_empresa=emp.id_empresa,
                cnpj=emp.cnpj,
                razao_social=emp.razao_social,
                tipo=emp.tipo
            ) for emp in empresas
        ]

# ==========================================
# SERVIÇO DE EQUIPAMENTO
# ==========================================
# Limites para não sobrecarregar o PartSurfer (endpoint não oficial da HP).
INTERVALO_MINIMO_POR_EQUIPAMENTO = 15   # segundos entre consultas do mesmo modelo
MAX_CONSULTAS_POR_MINUTO = 20           # total do sistema


class LimiteConsultasExcedido(Exception):
    """Muitas consultas ao PartSurfer em pouco tempo."""


def _controlar_taxa(id_equipamento: int) -> None:
    if not cache.add(f"partsurfer:equip:{id_equipamento}", 1, INTERVALO_MINIMO_POR_EQUIPAMENTO):
        raise LimiteConsultasExcedido(
            f"Este modelo foi consultado agora há pouco. Aguarde {INTERVALO_MINIMO_POR_EQUIPAMENTO} segundos."
        )
    chave = "partsurfer:global"
    cache.add(chave, 0, 60)
    try:
        total = cache.incr(chave)
    except ValueError:
        cache.set(chave, 1, 60)
        total = 1
    if total > MAX_CONSULTAS_POR_MINUTO:
        raise LimiteConsultasExcedido("Muitas consultas ao PartSurfer no último minuto. Tente novamente em instantes.")


class EquipamentoService:
    @staticmethod
    def salvar(dto: EquipamentoDTO) -> EquipamentoDTO:
        pn = (dto.pn or '').strip()
        modelo = (dto.modelo or '').strip()
        if not pn or not modelo:
            raise ValueError("Part Number (PN) e Modelo são obrigatórios.")
        dto.pn, dto.modelo = pn, modelo

        # UC12: impede cadastro duplicado no catálogo (o PN identifica a peça de forma única)
        duplicado = Equipamento.objects.filter(pn__iexact=pn)
        if dto.id_equipamento:
            duplicado = duplicado.exclude(pk=dto.id_equipamento)
        if duplicado.exists():
            raise ValueError("Já existe um equipamento cadastrado com este Part Number.")

        if dto.id_equipamento:
            equip = Equipamento.objects.get(pk=dto.id_equipamento)
            equip.pn = dto.pn
            equip.modelo = dto.modelo
            equip.save()
        else:
            equip = Equipamento.objects.create(
                pn=dto.pn,
                modelo=dto.modelo
            )
            
        return EquipamentoDTO(
            id_equipamento=equip.id_equipamento,
            pn=equip.pn,
            modelo=equip.modelo
        )

    @staticmethod
    def listar() -> List[EquipamentoDTO]:
        equipamentos = Equipamento.objects.annotate(total_pecas=Count('pecas'))
        return [
            EquipamentoDTO(
                id_equipamento=eq.id_equipamento,
                pn=eq.pn,
                modelo=eq.modelo,
                total_pecas=eq.total_pecas
            ) for eq in equipamentos
        ]

    # ---------- peças do modelo ----------
    @staticmethod
    def listar_pecas(id_equipamento: int) -> List[dict]:
        equip = Equipamento.objects.get(pk=id_equipamento)
        return [
            {'id': p.id_peca, 'part_number': p.part_number, 'descricao': p.descricao}
            for p in equip.pecas.order_by('part_number')
        ]

    @staticmethod
    def consultar_partsurfer(id_equipamento: int) -> List[dict]:
        """
        Consulta o HP PartSurfer pelo PN do modelo e devolve as peças CANDIDATAS, sem gravar nada.

        O usuário escolhe quais usar (ver salvar_pecas), porque nem tudo que o PartSurfer lista é útil.
        Peças já vinculadas vêm marcadas (selecionada=True); as vinculadas que a HP não retornou
        continuam na lista (somente_local=True) para o usuário decidir se mantém.

        Levanta PartSurferError (falha de comunicação), ValueError (HP sem peças) e LimiteConsultasExcedido.
        """
        equip = Equipamento.objects.get(pk=id_equipamento)
        _controlar_taxa(id_equipamento)
        pecas_hp = buscar_pecas_por_pn(equip.pn)

        if not pecas_hp:
            raise ValueError(
                f"O PartSurfer não retornou peças para o PN {equip.pn}. Confira se o Part Number está correto."
            )

        vinculadas = {p.part_number: p for p in equip.pecas.all()}
        candidatas, vistos = [], set()
        for item in pecas_hp:
            pn = str(item.get('part_number') or '').strip().upper()
            if not pn or pn in vistos:
                continue
            vistos.add(pn)
            candidatas.append({'part_number': pn, 'descricao': item.get('descricao', ''),
                               'selecionada': pn in vinculadas, 'somente_local': False})
        for pn, peca in vinculadas.items():
            if pn not in vistos:
                candidatas.append({'part_number': pn, 'descricao': peca.descricao,
                                   'selecionada': True, 'somente_local': True})
        candidatas.sort(key=lambda c: c['part_number'])
        return candidatas

    @staticmethod
    def salvar_pecas(id_equipamento: int, itens: List[dict]) -> List[dict]:
        """
        Define exatamente quais peças pertencem ao modelo: [{part_number, descricao}, ...].
        Peça nova entra no catálogo; peça removida só perde o vínculo com este modelo
        (continua no catálogo e no histórico das triagens).
        """
        if not isinstance(itens, list):
            raise ValueError("Lista de peças inválida.")
        equip = Equipamento.objects.get(pk=id_equipamento)

        escolhidas = {}
        for item in itens:
            if not isinstance(item, dict):
                continue
            pn = str(item.get('part_number') or '').strip().upper()
            if not pn:
                continue
            if len(pn) > 100:
                raise ValueError(f"Part Number muito longo: {pn[:20]}...")
            escolhidas[pn] = str(item.get('descricao') or '').strip()[:255]
        if len(escolhidas) > 200:
            raise ValueError("Limite de 200 peças por modelo.")

        with transaction.atomic():
            existentes = {p.part_number: p for p in Peca.objects.filter(part_number__in=list(escolhidas))}
            vinculos = []
            for pn, descricao in escolhidas.items():
                peca = existentes.get(pn)
                if peca is None:
                    peca = Peca.objects.create(part_number=pn, descricao=descricao)
                elif descricao and peca.descricao != descricao:
                    peca.descricao = descricao
                    peca.save(update_fields=['descricao'])
                vinculos.append(peca)
            equip.pecas.set(vinculos)
        return EquipamentoService.listar_pecas(id_equipamento)


# ==========================================
# SERVIÇO DE FUNCIONÁRIO
# ==========================================
class FuncionarioService:
    SENHA_MIN = 4

    # ---------- senhas ----------
    @staticmethod
    def _senha_em_hash(valor: str) -> bool:
        try:
            identify_hasher(valor)
            return True
        except ValueError:
            return False

    @staticmethod
    def _conferir_senha(func: Funcionario, senha_informada: str) -> bool:
        """Confere a senha. Registros antigos (texto puro) são aceitos uma vez e convertidos para hash."""
        armazenada = func.senha or ''
        if FuncionarioService._senha_em_hash(armazenada):
            return check_password(senha_informada, armazenada)
        if constant_time_compare(str(senha_informada), armazenada):
            func.senha = make_password(senha_informada)
            func.save(update_fields=['senha'])
            return True
        return False

    # ---------- cadastro ----------
    @staticmethod
    def salvar(dto: FuncionarioDTO) -> FuncionarioDTO:
        nome = (dto.nome or '').strip()
        if not nome:
            raise ValueError("O nome é obrigatório.")
        if not dto.cargo:
            raise ValueError("O cargo é obrigatório.")
        cargo = dto.cargo.value if hasattr(dto.cargo, 'value') else dto.cargo

        email = (dto.email or '').strip()
        if not email:
            raise ValueError("O e-mail é obrigatório.")
        try:
            validate_email(email)
        except ValidationError:
            raise ValueError("E-mail inválido.")

        duplicado = Funcionario.objects.filter(login=dto.login)
        if dto.id_funcionario:
            duplicado = duplicado.exclude(pk=dto.id_funcionario)
        if duplicado.exists():
            raise ValueError("Já existe um funcionário com esta matrícula.")

        if dto.id_funcionario:
            func = Funcionario.objects.get(pk=dto.id_funcionario)
            func.cargo = cargo
            func.nome = nome
            func.login = dto.login
            func.email = email
            if dto.senha:
                # Redefinição feita pelo Administrador (UC09 - FA02): exige nova troca no próximo login
                func.senha = make_password(dto.senha)
                func.primeiro_acesso = True
            func.save()
            if dto.senha:
                # Encerra os pedidos de recuperação pendentes deste funcionário
                SolicitacaoSenha.objects.filter(funcionario=func, atendida=False).update(atendida=True)
        else:
            # UC07: sem senha informada, a matrícula vira a senha provisória
            senha_provisoria = dto.senha or str(dto.login)
            func = Funcionario.objects.create(
                cargo=cargo,
                nome=nome,
                login=dto.login,
                email=email,
                senha=make_password(senha_provisoria),
                primeiro_acesso=True
            )

        return FuncionarioDTO(
            id_funcionario=func.id_funcionario,
            cargo=func.cargo,
            nome=func.nome,
            login=func.login,
            primeiro_acesso=func.primeiro_acesso,
            email=func.email
        )

    @staticmethod
    def listar() -> List[FuncionarioDTO]:
        funcionarios = Funcionario.objects.all()
        return [
            FuncionarioDTO(
                id_funcionario=func.id_funcionario,
                cargo=func.cargo,
                nome=func.nome,
                login=func.login,
                primeiro_acesso=func.primeiro_acesso,
                email=func.email
            ) for func in funcionarios
        ]

    # ---------- autenticação ----------
    @staticmethod
    def autenticar(login: int, senha: str) -> Optional[FuncionarioDTO]:
        func = Funcionario.objects.filter(login=login).first()

        if func and FuncionarioService._conferir_senha(func, senha):
            return FuncionarioDTO(
                id_funcionario=func.id_funcionario,
                cargo=func.cargo,
                nome=func.nome,
                login=func.login,
                primeiro_acesso=func.primeiro_acesso,
                email=func.email
            )
        return None

    @staticmethod
    def alterar_senha(id_funcionario: int, senha_atual: str, nova_senha: str) -> None:
        """RF13/RF14: troca de senha pelo próprio usuário (inclusive a troca obrigatória do primeiro acesso)."""
        func = Funcionario.objects.filter(pk=id_funcionario).first()
        if not func:
            raise ValueError("Funcionário não encontrado.")
        if not senha_atual or not nova_senha:
            raise ValueError("Informe a senha atual e a nova senha.")
        if not FuncionarioService._conferir_senha(func, senha_atual):
            raise ValueError("Senha atual incorreta.")
        if len(nova_senha) < FuncionarioService.SENHA_MIN:
            raise ValueError(f"A nova senha deve ter pelo menos {FuncionarioService.SENHA_MIN} caracteres.")
        if nova_senha == senha_atual:
            raise ValueError("A nova senha deve ser diferente da atual.")
        if nova_senha == str(func.login):
            raise ValueError("A nova senha não pode ser igual à matrícula.")

        func.senha = make_password(nova_senha)
        func.primeiro_acesso = False
        func.save(update_fields=['senha', 'primeiro_acesso'])

class SolicitacaoSenhaService:
    """UC09: o usuário informa matrícula e e-mail cadastrados; o sistema avisa o administrador por e-mail."""

    @staticmethod
    def registrar(login: int, email: str) -> None:
        email = (email or '').strip()
        if not email:
            return
        # Só aceita o pedido quando matrícula e e-mail conferem com o cadastro
        func = Funcionario.objects.filter(login=login, email__iexact=email).first()
        if not func:
            return
        # Um único pedido pendente por funcionário evita repetição e spam
        if SolicitacaoSenha.objects.filter(funcionario=func, atendida=False).exists():
            return

        pedido = SolicitacaoSenha.objects.create(funcionario=func)
        SolicitacaoSenhaService._avisar_administradores(func, pedido)

    @staticmethod
    def _destinatarios_admin() -> list:
        emails = list(
            Funcionario.objects.filter(cargo='Administrador').exclude(email='').values_list('email', flat=True)
        )
        extras = getattr(settings, 'DAEDALUS_EMAIL_ADMIN', '') or ''
        emails += [e.strip() for e in extras.split(',') if e.strip()]
        vistos, unicos = set(), []
        for e in emails:
            if e.lower() not in vistos:
                vistos.add(e.lower())
                unicos.append(e)
        return unicos

    @staticmethod
    def _avisar_administradores(func: Funcionario, pedido: SolicitacaoSenha) -> None:
        destinatarios = SolicitacaoSenhaService._destinatarios_admin()
        if not destinatarios:
            logger.warning("Pedido de redefinição de senha sem administrador com e-mail cadastrado (matrícula %s).", func.login)
            return

        data = timezone.localtime(pedido.data_solicitacao).strftime('%d/%m/%Y %H:%M')
        corpo = (
            f"O funcionário {func.nome} (matrícula {func.login}, e-mail {func.email}) "
            f"solicitou a redefinição de senha em {data}.\n\n"
            "Para atender: acesse Funcionário no Daedalus Gestio, localize o pedido no painel "
            "\"Solicitações de redefinição de senha\", clique em \"Redefinir senha\" e informe uma senha provisória. "
            "No próximo login o usuário será obrigado a trocá-la.\n\n"
            "Você pode responder a este e-mail para falar diretamente com o solicitante."
        )
        try:
            EmailMessage(
                subject=f"[Daedalus Gestio] Redefinição de senha solicitada: {func.nome}",
                body=corpo,
                from_email=settings.DEFAULT_FROM_EMAIL,
                to=destinatarios,
                reply_to=[func.email],
            ).send(fail_silently=False)
        except Exception:
            # O pedido continua visível no painel mesmo se o servidor de e-mail falhar
            logger.exception("Falha ao enviar o e-mail de redefinição de senha (matrícula %s).", func.login)

    @staticmethod
    def listar_pendentes() -> list:
        pedidos = SolicitacaoSenha.objects.filter(atendida=False).select_related('funcionario').order_by('data_solicitacao')
        return [
            {
                'id': p.id_solicitacao,
                'id_funcionario': p.funcionario.id_funcionario,
                'nome': p.funcionario.nome,
                'cargo': p.funcionario.cargo,
                'login': p.funcionario.login,
                'email': p.funcionario.email,
                'data': timezone.localtime(p.data_solicitacao).strftime('%d/%m/%Y %H:%M'),
            }
            for p in pedidos
        ]


# ==========================================
# SERVIÇO DE TÉCNICO
# ==========================================
class TecnicoService:
    @staticmethod
    def salvar(dto: TecnicoDTO) -> TecnicoDTO:
        # Busca a instância da Empresa no banco usando o ID recebido da tela
        empresa_obj = Empresa.objects.get(pk=dto.id_empresa)
        
        if dto.id_tecnico:
            tec = Tecnico.objects.get(pk=dto.id_tecnico)
            tec.nome = dto.nome
            tec.empresa = empresa_obj
            tec.save()
        else:
            tec = Tecnico.objects.create(nome=dto.nome, empresa=empresa_obj)
            
        return TecnicoDTO(
            id_tecnico=tec.id_tecnico,
            nome=tec.nome,
            id_empresa=tec.empresa.id_empresa
        )

    @staticmethod
    def listar() -> List[TecnicoDTO]:
        # select_related('empresa') evita o problema de N+1 queries no banco
        tecnicos = Tecnico.objects.select_related('empresa').all()
        return [
            TecnicoDTO(
                id_tecnico=tec.id_tecnico,
                nome=tec.nome,
                id_empresa=tec.empresa.id_empresa,
                nome_empresa=tec.empresa.razao_social
            ) for tec in tecnicos
        ]

# ==========================================
# SERVIÇO DE TRIAGEM
# ==========================================
class TriagemService:

    @staticmethod
    def validar(dto: TriagemDTO):
        if not dto.n_ticket:
            raise ValueError("O número do Ticket é obrigatório.")
        if not dto.sn or not str(dto.sn).strip():
            raise ValueError("O Serial Number (SN) é obrigatório.")
        # Validação atualizada para exigir os DOIS analistas
        if not dto.id_tecnico or not dto.id_cliente or not dto.id_analista_chamado or not dto.id_analista_triagem or not dto.id_equipamento:
            raise ValueError("Todos os vínculos (Técnico, Cliente, Analistas e Equipamento) são obrigatórios.")

    @staticmethod
    def salvar(dto: TriagemDTO) -> TriagemDTO:
        TriagemService.validar(dto)

        tecnico = Tecnico.objects.get(pk=dto.id_tecnico)
        cliente = Empresa.objects.get(pk=dto.id_cliente)
        equipamento = Equipamento.objects.get(pk=dto.id_equipamento)

        # Peça trocada (opcional): precisa pertencer ao modelo escolhido
        peca_trocada = None
        if dto.id_peca_trocada:
            peca_trocada = equipamento.pecas.filter(pk=dto.id_peca_trocada).first()
            if peca_trocada is None:
                raise ValueError("A peça trocada não pertence ao modelo de equipamento selecionado.")
        
        # Buscando as instâncias dos dois analistas
        analista_chamado = Funcionario.objects.get(pk=dto.id_analista_chamado)
        analista_triagem = Funcionario.objects.get(pk=dto.id_analista_triagem)

        if dto.id_triagem:
            triagem = Triagem.objects.get(pk=dto.id_triagem)
            triagem.n_ticket = dto.n_ticket
            triagem.nf = dto.nf
            triagem.cidade = dto.cidade
            triagem.estado = dto.estado
            triagem.sn = dto.sn
            triagem.status = dto.status_peca.value if hasattr(dto.status_peca, 'value') else dto.status_peca
            triagem.troca = dto.troca
            triagem.peca_trocada = peca_trocada
            triagem.defeito_usuario = dto.defeito_usuario
            triagem.defeito_tecnico = dto.defeito_tecnico
            triagem.defeito_constatado = dto.defeito_constatado
            triagem.testes = dto.testes
            triagem.analise = dto.analise
            triagem.conclusao = dto.conclusao
            
            triagem.tecnico = tecnico
            triagem.cliente = cliente
            triagem.equipamento = equipamento
            # Atualizando os dois analistas
            triagem.analista_chamado = analista_chamado
            triagem.analista_triagem = analista_triagem
            
            triagem.save()
        else:
            triagem = Triagem.objects.create(
                n_ticket=dto.n_ticket,
                nf=dto.nf,
                cidade=dto.cidade,
                estado=dto.estado,
                sn=dto.sn,
                status=dto.status_peca.value if hasattr(dto.status_peca, 'value') else dto.status_peca,
                troca=dto.troca,
                peca_trocada=peca_trocada,
                defeito_usuario=dto.defeito_usuario,
                defeito_tecnico=dto.defeito_tecnico,
                defeito_constatado=dto.defeito_constatado,
                testes=dto.testes,
                analise=dto.analise,
                conclusao=dto.conclusao,
                tecnico=tecnico,
                cliente=cliente,
                equipamento=equipamento,
                # Salvando os dois analistas
                analista_chamado=analista_chamado,
                analista_triagem=analista_triagem
            )

        dto.id_triagem = triagem.id_triagem
        return dto

    @staticmethod
    def listar() -> List[TriagemDTO]:
        triagens = Triagem.objects.select_related('cliente', 'equipamento', 'tecnico', 'analista_chamado', 'analista_triagem', 'peca_trocada').all()
        return [TriagemService._mapear_para_dto(t) for t in triagens]

    @staticmethod
    def buscar_por_id(id_triagem: int) -> Optional[TriagemDTO]:
        triagem = Triagem.objects.select_related('cliente', 'equipamento', 'tecnico', 'analista_chamado', 'analista_triagem', 'peca_trocada').filter(pk=id_triagem).first()
        if triagem:
            return TriagemService._mapear_para_dto(triagem)
        return None

    @staticmethod
    def buscar_por_ticket(ticket: int) -> Optional[TriagemDTO]:
        if not ticket:
            raise ValueError("Ticket obrigatório para a busca.")
            
        triagem = Triagem.objects.select_related('cliente', 'equipamento', 'tecnico', 'analista_chamado', 'analista_triagem', 'peca_trocada').filter(n_ticket=ticket).first()
        if triagem:
            return TriagemService._mapear_para_dto(triagem)
        return None

    @staticmethod
    def buscar_por_filtro(campo: str, valor: str) -> List[TriagemDTO]:
        if not valor or not campo:
            raise ValueError("Campo e valor de busca são obrigatórios.")
        
        campos_permitidos = ('n_ticket', 'nf', 'sn', 'cidade', 'estado',
                             'tecnico__nome', 'cliente__razao_social', 'equipamento__modelo')
        if campo not in campos_permitidos:
            raise ValueError("Campo de busca inválido.")

        filtro = {f"{campo}__icontains": valor}
        triagens = Triagem.objects.select_related('cliente', 'equipamento', 'tecnico', 'analista_chamado', 'analista_triagem', 'peca_trocada').filter(**filtro)
        return [TriagemService._mapear_para_dto(t) for t in triagens]

    @staticmethod
    def deletar(ticket: int):
        if not ticket:
            raise ValueError("Ticket obrigatório para exclusão.")
        
        for foto in FotoTriagem.objects.filter(triagem__n_ticket=ticket):
            foto.arquivo.delete(save=False)
        registros_apagados, _ = Triagem.objects.filter(n_ticket=ticket).delete()
        if registros_apagados == 0:
            raise ObjectDoesNotExist("Nenhuma triagem encontrada com este ticket.")
        

    # =======================================================
    # MÉTRICAS E ESTATÍSTICAS
    # =======================================================

    @staticmethod
    def total_triagens() -> int:
        return Triagem.objects.count()

    @staticmethod
    def total_peca_funcional() -> int:
        return Triagem.objects.filter(status=StatusPeca.FUNCIONAL).count()

    @staticmethod
    def total_peca_nao_funcional() -> int:
        return Triagem.objects.filter(status=StatusPeca.NAO_FUNCIONAL).count()

    # =======================================================
    # MÉTODOS AUXILIARES (Privados)
    # =======================================================
    
    @staticmethod
    def _mapear_para_dto(triagem: Triagem) -> TriagemDTO:
        return TriagemDTO(
            id_triagem=triagem.id_triagem,
            n_ticket=triagem.n_ticket,
            nf=triagem.nf,
            cidade=triagem.cidade,
            estado=triagem.estado,
            id_tecnico=triagem.tecnico_id,
            id_cliente=triagem.cliente_id,
            id_equipamento=triagem.equipamento_id,
            
            # --- Mapeamento dos DOIS Analistas ---
            id_analista_chamado=triagem.analista_chamado_id,
            id_analista_triagem=triagem.analista_triagem_id,
            
            sn=triagem.sn,
            status_peca=StatusPeca(triagem.status),
            troca=triagem.troca,
            id_peca_trocada=triagem.peca_trocada_id,
            peca_trocada=str(triagem.peca_trocada) if triagem.peca_trocada_id else None,
            defeito_usuario=triagem.defeito_usuario,
            defeito_tecnico=triagem.defeito_tecnico,
            defeito_constatado=triagem.defeito_constatado,
            testes=triagem.testes,
            analise=triagem.analise,
            conclusao=triagem.conclusao,
            data_triagem=triagem.data_triagem,
            
            # --- Nomes para o Front-end e Exportação ---
            nome_cliente=triagem.cliente.razao_social,
            nome_equipamento=triagem.equipamento.modelo,
            nome_tecnico=triagem.tecnico.nome,
            nome_analista_chamado=triagem.analista_chamado.nome,
            nome_analista_triagem=triagem.analista_triagem.nome
        )

    @staticmethod
    def gerar_docx(id_triagem: int) -> tuple[BytesIO, str]:
        tri = TriagemService.buscar_por_id(id_triagem)
        if not tri:
            raise ValueError("Triagem não encontrada")

        template_path = os.path.join(settings.BASE_DIR, 'templates', 'Daedalus', 'TRIAGEM.docx')
        
        if not os.path.exists(template_path):
            raise FileNotFoundError("Modelo TRIAGEM.docx não encontrado.")
            
        doc = Document(template_path)
        # Data real da triagem (e não a data da exportação)
        data_triagem = timezone.localtime(tri.data_triagem).strftime("%d/%m/%Y") if tri.data_triagem else ""

        # ==========================================
        # TABELA 2: DADOS DO CHAMADO
        # ==========================================
        t2 = doc.tables[2]
        t2.cell(0, 0).text = f"Data: {data_triagem}"
        t2.cell(0, 1).text = f"NF: {tri.nf or ''}"
        t2.cell(0, 2).text = f"Ticket: {tri.n_ticket or ''}"
        
        t2.cell(1, 1).text = f"Técnico: {tri.nome_tecnico or ''}"
        
        # Insere ambos os analistas na tabela
        t2.cell(1, 2).text = f"Analista Chamado: {tri.nome_analista_chamado or ''} | Triagem: {tri.nome_analista_triagem or ''}"
        
        t2.cell(2, 1).text = f"Região: {tri.cidade or ''} - {tri.estado or ''}"
        t2.cell(2, 2).text = f"Equipamento: {tri.nome_equipamento or ''}"
        
        t2.cell(3, 1).text = f"N° de Série: {tri.sn or ''}"
        trocado = ' | '.join(x for x in (tri.peca_trocada, tri.troca) if x)
        t2.cell(4, 1).text = f"Peça/Equipamento Trocado: {trocado}"
        t2.cell(5, 1).text = f"Defeito alegado: {tri.defeito_usuario or ''}"
        t2.cell(6, 1).text = f"Constatação do Técnico: {tri.defeito_tecnico or ''}"

        # ==========================================
        # TABELA 4: LAUDO TÉCNICO
        # ==========================================
        t4 = doc.tables[4]
        t4.cell(0, 0).text = f"Defeito constatado: {tri.defeito_constatado or ''}"
        t4.cell(1, 0).text = f"Testes Realizados: {tri.testes or ''}"
        t4.cell(2, 0).text = f"Análise: {tri.analise or ''}"
        t4.cell(3, 0).text = f"Conclusão: {tri.conclusao or ''}"

        FotoService.inserir_no_docx(doc, id_triagem)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        # ==========================================
        # NOME DO ARQUIVO FORMATADO
        # ==========================================
        modelo_str = str(tri.nome_equipamento or 'SemModelo').replace(" ", "")
        cliente_str = str(tri.nome_cliente or 'SemCliente').replace(" ", "")
        ticket_str = str(tri.n_ticket)
        
        nome_arquivo = f"Triagem_{ticket_str}_{modelo_str}_{cliente_str}.docx"
        
        return buffer, nome_arquivo

# ==========================================
# SERVIÇO DE FOTOS DA TRIAGEM
# ==========================================
class FotoService:
    MAX_FOTOS = 10
    MAX_BYTES = 10 * 1024 * 1024          # 10 MB por arquivo enviado
    MAX_LADO = 1600                       # px; fotos maiores são reduzidas
    FORMATOS = ('JPEG', 'PNG', 'WEBP')

    @staticmethod
    def _para_dict(foto: FotoTriagem) -> dict:
        return {
            'id': foto.id_foto,
            'url': f"/fotos/{foto.id_foto}/",
            'legenda': foto.legenda,
            'criado_em': timezone.localtime(foto.criado_em).isoformat(),
        }

    @staticmethod
    def listar(id_triagem: int) -> List[dict]:
        return [FotoService._para_dict(f) for f in FotoTriagem.objects.filter(triagem_id=id_triagem)]

    @staticmethod
    def _processar(arquivo) -> bytes:
        """Valida que é uma imagem de verdade, corrige a rotação, remove metadados (EXIF/GPS) e reduz o tamanho."""
        if arquivo.size > FotoService.MAX_BYTES:
            raise ValueError("A foto excede o limite de 10 MB.")
        try:
            img = Image.open(arquivo)
            if img.format not in FotoService.FORMATOS:
                raise ValueError("Formato não suportado. Envie JPG, PNG ou WEBP.")
            img = ImageOps.exif_transpose(img)
            img.load()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError):
            raise ValueError("O arquivo enviado não é uma imagem válida.")
        if img.mode not in ('RGB', 'L'):
            fundo = Image.new('RGB', img.size, (255, 255, 255))
            rgba = img.convert('RGBA')
            fundo.paste(rgba, mask=rgba.split()[-1])
            img = fundo
        elif img.mode == 'L':
            img = img.convert('RGB')
        img.thumbnail((FotoService.MAX_LADO, FotoService.MAX_LADO))
        saida = BytesIO()
        img.save(saida, format='JPEG', quality=85, optimize=True)
        return saida.getvalue()

    @staticmethod
    def adicionar(id_triagem: int, arquivo, legenda: str = '') -> dict:
        triagem = Triagem.objects.get(pk=id_triagem)
        if triagem.fotos.count() >= FotoService.MAX_FOTOS:
            raise ValueError(f"Limite de {FotoService.MAX_FOTOS} fotos por triagem.")
        conteudo = FotoService._processar(arquivo)
        foto = FotoTriagem(triagem=triagem, legenda=(legenda or '').strip()[:150])
        foto.arquivo.save(f"{uuid.uuid4().hex}.jpg", ContentFile(conteudo), save=False)
        foto.save()
        return FotoService._para_dict(foto)

    @staticmethod
    def excluir(id_foto: int) -> None:
        foto = FotoTriagem.objects.filter(pk=id_foto).first()
        if foto is None:
            raise ObjectDoesNotExist("Foto não encontrada.")
        foto.arquivo.delete(save=False)
        foto.delete()

    @staticmethod
    def inserir_no_docx(doc, id_triagem: int) -> None:
        """Acrescenta o registro fotográfico ao final do laudo (2 fotos por linha)."""
        from docx.shared import Cm
        fotos = list(FotoTriagem.objects.filter(triagem_id=id_triagem))
        if not fotos:
            return
        doc.add_paragraph().add_run("Registro fotográfico").bold = True
        linhas = (len(fotos) + 1) // 2
        tabela = doc.add_table(rows=linhas, cols=2)
        for i, foto in enumerate(fotos):
            celula = tabela.cell(i // 2, i % 2)
            par = celula.paragraphs[0]
            try:
                with foto.arquivo.open('rb') as f:
                    par.add_run().add_picture(BytesIO(f.read()), width=Cm(7.0))
            except (FileNotFoundError, OSError):
                par.add_run("[foto indisponível]")
            if foto.legenda:
                celula.add_paragraph(foto.legenda)
