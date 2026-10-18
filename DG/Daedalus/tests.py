import json
import shutil
import tempfile
import time
from io import BytesIO
from unittest import mock
from datetime import datetime

from django.contrib.auth.hashers import make_password
from django.core import mail
from django.test import Client, TestCase, override_settings
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document
from PIL import Image

from .analytics import classificar, intervalo_wilson, normalizar_texto
from .dtos import CargoFuncionario, EquipamentoDTO, FuncionarioDTO, StatusPeca, TriagemDTO
from .models import Empresa, Equipamento, FotoTriagem, Funcionario, Peca, SolicitacaoSenha, Tecnico, Triagem
from . import partsurfer
from .services import EquipamentoService, FotoService, FuncionarioService, TriagemService


# ------------------------------------------------------------------
# Auxiliares
# ------------------------------------------------------------------
def criar_funcionario(login, cargo='Analista de Triagem', senha='senha123', primeiro_acesso=False, email=None):
    return Funcionario.objects.create(
        cargo=cargo, nome=f'Funcionario {login}', login=login,
        senha=make_password(senha), primeiro_acesso=primeiro_acesso,
        email=email if email is not None else f'func{login}@empresa.com'
    )


def criar_triagem(analista, n_ticket=1, data=None, status=1, **extra):
    prestadora, _ = Empresa.objects.get_or_create(
        cnpj='11111111000111', defaults={'razao_social': 'Prestadora X', 'tipo': 1})
    cliente, _ = Empresa.objects.get_or_create(
        cnpj='22222222000122', defaults={'razao_social': 'Cliente Y', 'tipo': 2})
    tecnico, _ = Tecnico.objects.get_or_create(nome='Tecnico Z', empresa=prestadora)
    equip, _ = Equipamento.objects.get_or_create(pn='PN-1', defaults={'modelo': 'Notebook A'})
    dados = dict(
        n_ticket=n_ticket, nf=10, cidade='Sorocaba', estado='SP',
        tecnico=tecnico, cliente=cliente, equipamento=equip,
        analista_chamado=analista, analista_triagem=analista,
        sn=f'SN{n_ticket}', status=status, troca='Placa-mae',
        defeito_usuario='Nao liga', defeito_tecnico='Placa-mae',
        defeito_constatado='Bateria', testes='Teste de bancada',
        analise='Analise', conclusao='Conclusao',
    )
    if data:
        dados['data_triagem'] = data
    dados.update(extra)
    return Triagem.objects.create(**dados)


def post_json(client, url, corpo):
    return client.post(url, json.dumps(corpo), content_type='application/json')


def logar(client, login, senha='senha123'):
    resposta = post_json(client, '/api/login/', {'login': login, 'senha': senha})
    assert resposta.status_code == 200, resposta.content
    return resposta


# ------------------------------------------------------------------
# Autenticação, senhas e perfis (RF10, RF12, RF13, RF14, RNF05)
# ------------------------------------------------------------------
class AutenticacaoTests(TestCase):

    def test_rotas_exigem_login(self):
        resposta = self.client.get('/principal/')
        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(resposta.url, '/')
        self.assertEqual(self.client.get('/api/triagem/').status_code, 401)
        self.assertEqual(self.client.get('/api/dashboard/').status_code, 401)

    def test_login_valido_e_invalido(self):
        criar_funcionario(100)
        ruim = post_json(self.client, '/api/login/', {'login': 100, 'senha': 'errada'})
        self.assertEqual(ruim.status_code, 401)
        logar(self.client, 100)
        self.assertEqual(self.client.session['funcionario_id'], Funcionario.objects.get(login=100).id_funcionario)

    def test_senha_nao_fica_em_texto_puro(self):
        FuncionarioService.salvar(FuncionarioDTO(
            cargo=CargoFuncionario.ANALISTA_TRIAGEM, nome='Ana', login=200, senha='abc12345',
            email='ana@empresa.com'))
        senha_salva = Funcionario.objects.get(login=200).senha
        self.assertNotEqual(senha_salva, 'abc12345')
        self.assertTrue(FuncionarioService._senha_em_hash(senha_salva))

    def test_senha_legada_em_texto_puro_e_convertida_no_login(self):
        Funcionario.objects.create(cargo='Gestor', nome='Antigo', login=210, senha='1234', primeiro_acesso=False)
        logar(self.client, 210, '1234')
        self.assertTrue(FuncionarioService._senha_em_hash(Funcionario.objects.get(login=210).senha))

    def test_matricula_vira_senha_provisoria(self):
        FuncionarioService.salvar(FuncionarioDTO(
            cargo=CargoFuncionario.FUNCIONARIO, nome='Novo', login=300, email='novo@empresa.com'))
        dto = FuncionarioService.autenticar(300, '300')
        self.assertIsNotNone(dto)
        self.assertTrue(dto.primeiro_acesso)

    def test_email_obrigatorio_e_valido(self):
        for email in ('', '   ', 'sem-arroba'):
            with self.assertRaises(ValueError):
                FuncionarioService.salvar(FuncionarioDTO(
                    cargo=CargoFuncionario.FUNCIONARIO, nome='Sem Email', login=320, email=email))
        self.assertFalse(Funcionario.objects.filter(login=320).exists())

    def test_matricula_duplicada_e_rejeitada(self):
        criar_funcionario(310)
        with self.assertRaises(ValueError):
            FuncionarioService.salvar(FuncionarioDTO(
                cargo=CargoFuncionario.FUNCIONARIO, nome='Outro', login=310, email='outro@empresa.com'))

    def test_primeiro_acesso_bloqueia_telas_ate_trocar_a_senha(self):
        criar_funcionario(400, senha='400', primeiro_acesso=True)
        resposta = logar(self.client, 400, '400')
        self.assertTrue(resposta.json()['primeiro_acesso'])

        bloqueada = self.client.get('/principal/')
        self.assertEqual(bloqueada.status_code, 302)
        self.assertEqual(bloqueada.url, '/primeiro-acesso/')
        self.assertEqual(self.client.get('/api/triagem/').status_code, 403)

        errada = post_json(self.client, '/api/alterar-senha/', {'senha_atual': 'x', 'nova_senha': 'nova1234'})
        self.assertEqual(errada.status_code, 400)

        ok = post_json(self.client, '/api/alterar-senha/', {'senha_atual': '400', 'nova_senha': 'nova1234'})
        self.assertEqual(ok.status_code, 200)
        self.assertFalse(Funcionario.objects.get(login=400).primeiro_acesso)
        self.assertEqual(self.client.get('/api/triagem/').status_code, 200)

        # a senha nova passa a valer
        self.client.get('/logout/')
        logar(self.client, 400, 'nova1234')

    def test_regras_da_nova_senha(self):
        func = criar_funcionario(4100, senha='atual123')
        with self.assertRaises(ValueError):
            FuncionarioService.alterar_senha(func.id_funcionario, 'atual123', 'abc')        # curta
        with self.assertRaises(ValueError):
            FuncionarioService.alterar_senha(func.id_funcionario, 'atual123', 'atual123')   # igual a atual
        with self.assertRaises(ValueError):
            FuncionarioService.alterar_senha(func.id_funcionario, 'atual123', '4100')       # igual a matricula
        with self.assertRaises(ValueError):
            FuncionarioService.alterar_senha(func.id_funcionario, 'errada', 'nova12345')    # atual incorreta

    def test_reset_pelo_administrador_exige_nova_troca(self):
        func = criar_funcionario(420, primeiro_acesso=False)
        FuncionarioService.salvar(FuncionarioDTO(
            id_funcionario=func.id_funcionario, cargo=CargoFuncionario.FUNCIONARIO,
            nome=func.nome, login=420, senha='provisoria', email=func.email))
        func.refresh_from_db()
        self.assertTrue(func.primeiro_acesso)

    def test_somente_administrador_ou_gestor_gerencia_funcionarios(self):
        criar_funcionario(500, cargo='Analista de Triagem')
        criar_funcionario(501, cargo='Administrador')

        logar(self.client, 500)
        self.assertEqual(self.client.get('/api/funcionario/').status_code, 200)   # listagem liberada
        corpo = {'nome': 'Teste', 'cargo': 'Funcionário', 'login': 502, 'email': 'teste@empresa.com'}
        self.assertEqual(post_json(self.client, '/api/funcionario/', corpo).status_code, 403)
        self.assertEqual(self.client.get('/gerenciar-funcionario/').status_code, 302)
        self.client.get('/logout/')

        logar(self.client, 501)
        self.assertEqual(post_json(self.client, '/api/funcionario/', corpo).status_code, 200)
        self.assertEqual(self.client.get('/gerenciar-funcionario/').status_code, 200)


# ------------------------------------------------------------------
# Catálogo de equipamentos (UC12)
# ------------------------------------------------------------------
class EquipamentoTests(TestCase):

    def test_pn_duplicado_e_rejeitado(self):
        EquipamentoService.salvar(EquipamentoDTO(pn='ABC-1', modelo='Modelo 1'))
        with self.assertRaises(ValueError):
            EquipamentoService.salvar(EquipamentoDTO(pn='abc-1', modelo='Modelo 2'))

    def test_editar_o_proprio_registro_e_permitido(self):
        criado = EquipamentoService.salvar(EquipamentoDTO(pn='ABC-2', modelo='Modelo 1'))
        EquipamentoService.salvar(EquipamentoDTO(
            id_equipamento=criado.id_equipamento, pn='ABC-2', modelo='Modelo 1 revisado'))
        self.assertEqual(Equipamento.objects.get(pk=criado.id_equipamento).modelo, 'Modelo 1 revisado')

    def test_campos_obrigatorios(self):
        with self.assertRaises(ValueError):
            EquipamentoService.salvar(EquipamentoDTO(pn=' ', modelo='Modelo'))


# ------------------------------------------------------------------
# Dashboard / BI (RF05, UC03): filtros, comparação, controle estatístico e insights
# ------------------------------------------------------------------
def em(ano, mes, dia):
    return timezone.make_aware(datetime(ano, mes, dia, 12, 0))


class EstatisticaTests(TestCase):

    def test_intervalo_de_wilson(self):
        self.assertEqual(intervalo_wilson(0, 0), (0.0, 0.0))
        inf, sup = intervalo_wilson(5, 10)
        self.assertAlmostEqual(inf, 0.2366, places=3)
        self.assertAlmostEqual(sup, 0.7634, places=3)
        # poucos casos geram intervalo largo: 10 de 10 não vira "100% certo"
        self.assertLess(intervalo_wilson(10, 10)[0], 0.75)

    def test_classificacao_pela_carta_de_controle(self):
        self.assertEqual(classificar(0.9, 0.3, 4)[0], 'insuficiente')
        self.assertEqual(classificar(0.60, 0.3, 30)[0], 'critico')
        self.assertEqual(classificar(0.50, 0.3, 30)[0], 'atencao')
        self.assertEqual(classificar(0.45, 0.3, 30)[0], 'normal')
        self.assertEqual(classificar(0.10, 0.3, 30)[0], 'destaque')

    def test_normalizacao_de_texto(self):
        self.assertEqual(normalizar_texto('  Placa-mãe '), 'placa mae')
        self.assertEqual(normalizar_texto(None), '')


class DashboardTests(TestCase):

    def setUp(self):
        self.func = criar_funcionario(600, cargo='Gestor')
        criar_triagem(self.func, n_ticket=1, data=em(2026, 1, 10), status=1)
        criar_triagem(self.func, n_ticket=2, data=em(2026, 3, 10), status=2)
        logar(self.client, 600)

    def dados(self, query=''):
        resposta = self.client.get(f'/api/dashboard/{query}')
        self.assertEqual(resposta.status_code, 200)
        return resposta.json()['dados']

    def test_sem_filtro_traz_tudo(self):
        dados = self.dados()
        self.assertEqual(dados['kpis']['triagens']['valor'], 2)
        self.assertEqual(dados['kpis']['taxa_nff']['valor'], 50.0)

    def test_filtro_por_periodo(self):
        dados = self.dados('?inicio=2026-03-01&fim=2026-03-31')
        self.assertEqual(dados['kpis']['triagens']['valor'], 1)
        self.assertEqual(dados['kpis']['defeitos'], 1)
        self.assertEqual(dados['kpis']['taxa_nff']['valor'], 0.0)

    def test_periodo_sem_registros_devolve_opcoes_para_os_filtros(self):
        resposta = self.client.get('/api/dashboard/?inicio=2025-01-01&fim=2025-01-31').json()
        self.assertTrue(resposta['vazio'])
        self.assertIn('tecnicos', resposta['dados']['opcoes'])

    def test_datas_invalidas(self):
        self.assertEqual(self.client.get('/api/dashboard/?inicio=31-01-2026').status_code, 400)
        self.assertEqual(self.client.get('/api/dashboard/?inicio=2026-03-31&fim=2026-03-01').status_code, 400)
        self.assertEqual(self.client.get('/api/dashboard/?tecnico=abc').status_code, 400)

    def test_comparacao_com_o_periodo_anterior(self):
        criar_triagem(self.func, n_ticket=3, data=em(2026, 2, 10), status=1)   # cai no período anterior a março
        dados = self.dados('?inicio=2026-03-01&fim=2026-03-31')
        self.assertEqual(dados['meta']['anterior'], {'inicio': '2026-01-29', 'fim': '2026-02-28'})
        taxa = dados['kpis']['taxa_nff']
        self.assertEqual(taxa['anterior'], 100.0)
        self.assertEqual(taxa['delta'], -100.0)
        # desligando a comparação, o anterior some
        sem = self.dados('?inicio=2026-03-01&fim=2026-03-31&comparar=0')
        self.assertIsNone(sem['kpis']['taxa_nff']['anterior'])

    def test_serie_temporal_preenche_periodos_sem_dados(self):
        dados = self.dados('?inicio=2026-03-01&fim=2026-03-10')
        pontos = dados['serie']['pontos']
        self.assertEqual(dados['serie']['granularidade'], 'dia')
        self.assertEqual(len(pontos), 10)
        self.assertEqual(sum(p['total'] for p in pontos), 1)
        self.assertIsNone(pontos[0]['taxa_nff'])          # dia sem triagens não vira 0%

    def test_filtro_por_tecnico_e_prestadora(self):
        prestadora = Empresa.objects.get(cnpj='11111111000111')
        outro = Tecnico.objects.create(nome='Outro Tecnico', empresa=prestadora)
        criar_triagem(self.func, n_ticket=9, tecnico=outro, data=em(2026, 3, 11))
        self.assertEqual(self.dados(f'?tecnico={outro.id_tecnico}')['kpis']['triagens']['valor'], 1)
        self.assertEqual(self.dados(f'?prestadora={prestadora.id_empresa}')['kpis']['triagens']['valor'], 3)
        self.assertEqual(self.dados('?estado=sp')['kpis']['triagens']['valor'], 3)
        self.assertEqual(self.client.get('/api/dashboard/?estado=RJ').json()['vazio'], True)

    def test_acertos_erros_do_tecnico_e_divergencia_de_diagnostico(self):
        # peça com defeito confirmado, diagnóstico igual após normalizar (acento, caixa e pontuação)
        criar_triagem(self.func, n_ticket=3, status=2, data=em(2026, 3, 12),
                      defeito_tecnico='Placa-mãe', defeito_constatado='placa mae')
        dados = self.dados()
        tec = dados['tecnicos'][0]
        # acerto = defeito confirmado (status 2); erro = peça funcional (status 1)
        self.assertEqual((tec['acertos'], tec['nff'], tec['total']), (2, 1, 3))
        self.assertEqual(tec['taxa_acerto'], 66.7)
        self.assertEqual(tec['divergencias'], 2)      # tickets 1 e 2 divergem; o 3 coincide
        self.assertEqual(dados['kpis']['divergentes'], 2)
        self.assertEqual(dados['kpis']['comparaveis'], 3)

    def test_tecnico_fora_do_controle_gera_alerta_e_insight(self):
        prestadora = Empresa.objects.get(cnpj='11111111000111')
        ruim = Tecnico.objects.create(nome='Tecnico Ruim', empresa=prestadora)
        bom = Tecnico.objects.create(nome='Tecnico Bom', empresa=prestadora)
        Triagem.objects.all().delete()
        for i in range(10):
            criar_triagem(self.func, n_ticket=100 + i, tecnico=ruim, status=1, data=em(2026, 3, 5))
            criar_triagem(self.func, n_ticket=200 + i, tecnico=bom, status=2, data=em(2026, 3, 6))
        dados = self.dados()
        por_nome = {t['nome']: t for t in dados['tecnicos']}
        self.assertEqual(por_nome['Tecnico Ruim']['status'], 'critico')
        self.assertEqual(por_nome['Tecnico Bom']['status'], 'destaque')
        self.assertEqual(dados['kpis']['tecnicos_alerta']['criticos'], 1)
        titulos = [i['titulo'] for i in dados['insights']]
        self.assertTrue(any('Tecnico Ruim fora do limite de controle' in t for t in titulos))
        insight = next(i for i in dados['insights'] if 'Tecnico Ruim' in i['titulo'])
        self.assertEqual(insight['filtro'], {'tecnico': ruim.id_tecnico})   # permite drill-down
        self.assertEqual(dados['insights'][0]['nivel'], 'critico')     # o mais grave vem primeiro

    def test_amostra_pequena_nao_e_classificada(self):
        for i in range(2):                                   # 2 do setUp + 2 = 4 triagens (< 5)
            criar_triagem(self.func, n_ticket=300 + i, status=1, data=em(2026, 3, 7))
        tecnicos = self.dados()['tecnicos']
        self.assertEqual(tecnicos[0]['total'], 4)
        self.assertEqual(tecnicos[0]['status'], 'insuficiente')
        self.assertEqual(self.dados()['kpis']['tecnicos_alerta']['avaliados'], 0)

    def test_pareto_de_defeitos_agrupa_variacoes_de_escrita(self):
        Triagem.objects.all().delete()
        criar_triagem(self.func, n_ticket=11, status=2, defeito_constatado='Bateria')
        criar_triagem(self.func, n_ticket=12, status=2, defeito_constatado=' bateria ')
        criar_triagem(self.func, n_ticket=13, status=2, defeito_constatado='Tela')
        pareto = self.dados()['pareto_defeitos']
        self.assertEqual(pareto['total'], 3)
        self.assertEqual(pareto['itens'][0]['qtd'], 2)
        self.assertEqual(pareto['itens'][-1]['acumulado_pct'], 100.0)

    def test_reincidencia_por_numero_de_serie(self):
        criar_triagem(self.func, n_ticket=21, sn='abc123', data=em(2026, 3, 8))
        criar_triagem(self.func, n_ticket=22, sn=' ABC123 ', data=em(2026, 3, 9))
        reinc = self.dados()['reincidencia']
        self.assertEqual(reinc['series'], 1)
        self.assertEqual(reinc['triagens'], 2)
        self.assertEqual(reinc['itens'][0]['sn'], 'ABC123')
        self.assertEqual(reinc['itens'][0]['vezes'], 2)

    def test_diagnosticos_trocados_mais_frequentes(self):
        Triagem.objects.all().delete()
        for i in range(3):
            criar_triagem(self.func, n_ticket=30 + i, defeito_tecnico='Placa-mae', defeito_constatado='Bateria')
        criar_triagem(self.func, n_ticket=40, defeito_tecnico='Tela', defeito_constatado='Tela')
        trans = self.dados()['transicoes']
        self.assertEqual(trans['divergentes'], 3)
        self.assertEqual(trans['comparaveis'], 4)
        self.assertEqual(trans['itens'][0]['qtd'], 3)
        self.assertEqual(trans['itens'][0]['pct'], 100.0)

    def test_exportacao_csv_respeita_filtros_e_neutraliza_formulas(self):
        criar_triagem(self.func, n_ticket=50, data=em(2026, 3, 15), cidade='=SOMA(1;1)')
        resposta = self.client.get('/api/dashboard/exportar/?inicio=2026-03-01&fim=2026-03-31')
        self.assertEqual(resposta.status_code, 200)
        self.assertIn('text/csv', resposta['Content-Type'])
        texto = resposta.content.decode('utf-8-sig')
        linhas = texto.strip().splitlines()
        self.assertEqual(len(linhas), 3)               # cabeçalho + 2 triagens de março
        self.assertIn("'=SOMA(1;1)", texto)            # fórmula neutralizada
        self.assertNotIn(';=SOMA', texto)

    def test_exportacao_exige_login(self):
        self.client.get('/logout/')
        self.assertEqual(self.client.get('/api/dashboard/exportar/').status_code, 401)


# ------------------------------------------------------------------
# Triagem, laudo .docx e desempenho (RF02, RF06, RNF03, RNF04)
# ------------------------------------------------------------------
class TriagemELaudoTests(TestCase):

    def setUp(self):
        self.func = criar_funcionario(700)

    def test_laudo_usa_a_data_da_triagem(self):
        tri = criar_triagem(self.func, data=timezone.make_aware(datetime(2026, 1, 10, 12, 0)))
        buffer, nome = TriagemService.gerar_docx(tri.id_triagem)
        self.assertTrue(nome.startswith('Triagem_1_'))
        doc = Document(buffer)
        self.assertEqual(doc.tables[2].cell(0, 0).text, 'Data: 10/01/2026')

    def test_tempo_de_gravacao_menor_que_2_segundos(self):  # RNF03
        criar_triagem(self.func, n_ticket=1)
        base = Triagem.objects.get(n_ticket=1)
        dto = TriagemDTO(
            n_ticket=999, nf=1, cidade='Sorocaba', estado='SP',
            id_tecnico=base.tecnico_id, id_cliente=base.cliente_id,
            id_analista_chamado=self.func.id_funcionario,
            id_analista_triagem=self.func.id_funcionario,
            id_equipamento=base.equipamento_id, sn='SN999',
            status_peca=StatusPeca.FUNCIONAL,
        )
        inicio = time.perf_counter()
        TriagemService.salvar(dto)
        self.assertLess(time.perf_counter() - inicio, 2.0)

    def test_tempo_de_geracao_do_laudo_menor_que_5_segundos(self):  # RNF04
        tri = criar_triagem(self.func)
        inicio = time.perf_counter()
        TriagemService.gerar_docx(tri.id_triagem)
        self.assertLess(time.perf_counter() - inicio, 5.0)

    def test_filtro_de_busca_so_aceita_campos_permitidos(self):
        with self.assertRaises(ValueError):
            TriagemService.buscar_por_filtro('analista_triagem__senha', 'x')


# ------------------------------------------------------------------
# CSRF (agora ativo em todas as APIs)
# ------------------------------------------------------------------
class CsrfTests(TestCase):

    def test_post_sem_token_e_rejeitado(self):
        criar_funcionario(800)
        cliente = Client(enforce_csrf_checks=True)
        resposta = cliente.post('/api/login/', json.dumps({'login': 800, 'senha': 'senha123'}),
                                content_type='application/json')
        self.assertEqual(resposta.status_code, 403)

    def test_post_com_token_do_cookie_funciona(self):
        criar_funcionario(801)
        cliente = Client(enforce_csrf_checks=True)
        cliente.get('/')  # a tela de login entrega o cookie csrftoken
        token = cliente.cookies['csrftoken'].value
        resposta = cliente.post('/api/login/', json.dumps({'login': 801, 'senha': 'senha123'}),
                                content_type='application/json', HTTP_X_CSRFTOKEN=token)
        self.assertEqual(resposta.status_code, 200)


# ------------------------------------------------------------------
# Recuperação de senha por solicitação ao administrador (UC09)
# ------------------------------------------------------------------
class RecuperacaoSenhaTests(TestCase):

    def test_fluxo_completo_com_email_ao_administrador(self):
        func = criar_funcionario(900, senha='antiga123')
        criar_funcionario(901, cargo='Administrador')
        pedido = {'login': 900, 'email': 'FUNC900@empresa.com'}   # e-mail sem diferenciar maiúsculas

        # público: não exige login
        self.assertEqual(post_json(self.client, '/api/recuperar-senha/', pedido).status_code, 200)
        self.assertEqual(SolicitacaoSenha.objects.filter(funcionario=func, atendida=False).count(), 1)

        # o administrador recebe o e-mail e pode responder ao solicitante
        self.assertEqual(len(mail.outbox), 1)
        mensagem = mail.outbox[0]
        self.assertEqual(mensagem.to, ['func901@empresa.com'])
        self.assertEqual(mensagem.reply_to, ['func900@empresa.com'])
        self.assertIn('900', mensagem.body)

        # repetir o pedido não duplica nem reenvia o e-mail
        post_json(self.client, '/api/recuperar-senha/', pedido)
        self.assertEqual(SolicitacaoSenha.objects.count(), 1)
        self.assertEqual(len(mail.outbox), 1)

        # a lista de pedidos é restrita
        self.assertEqual(self.client.get('/api/solicitacoes-senha/').status_code, 401)
        logar(self.client, 901)
        pedidos = self.client.get('/api/solicitacoes-senha/').json()['dados']
        self.assertEqual(len(pedidos), 1)
        self.assertEqual(pedidos[0]['login'], 900)

        # o administrador redefine: pedido atendido e o usuário volta ao primeiro acesso
        corpo = {'id': func.id_funcionario, 'nome': func.nome, 'cargo': 'Analista de Triagem',
                 'login': 900, 'email': func.email, 'senha': 'provisoria1'}
        self.assertEqual(post_json(self.client, '/api/funcionario/', corpo).status_code, 200)
        self.assertEqual(len(self.client.get('/api/solicitacoes-senha/').json()['dados']), 0)
        func.refresh_from_db()
        self.assertTrue(func.primeiro_acesso)

    def test_dados_que_nao_conferem_nao_geram_pedido_nem_email(self):
        criar_funcionario(910)
        criar_funcionario(911, cargo='Administrador')
        casos = [
            {'login': 910, 'email': 'outro@empresa.com'},   # e-mail errado
            {'login': 99999, 'email': 'func910@empresa.com'},  # matrícula inexistente
            {'login': 910},                                   # sem e-mail
        ]
        for caso in casos:
            self.assertEqual(post_json(self.client, '/api/recuperar-senha/', caso).status_code, 200)
        self.assertEqual(SolicitacaoSenha.objects.count(), 0)
        self.assertEqual(len(mail.outbox), 0)

    def test_sem_administrador_com_email_o_pedido_fica_no_painel(self):
        criar_funcionario(920)
        criar_funcionario(921, cargo='Administrador', email='')
        post_json(self.client, '/api/recuperar-senha/', {'login': 920, 'email': 'func920@empresa.com'})
        self.assertEqual(SolicitacaoSenha.objects.count(), 1)
        self.assertEqual(len(mail.outbox), 0)

    @override_settings(DAEDALUS_EMAIL_ADMIN='ti@empresa.com, suporte@empresa.com')
    def test_destinatarios_extras_do_ambiente(self):
        criar_funcionario(930)
        criar_funcionario(931, cargo='Administrador')
        post_json(self.client, '/api/recuperar-senha/', {'login': 930, 'email': 'func930@empresa.com'})
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(sorted(mail.outbox[0].to), ['func931@empresa.com', 'suporte@empresa.com', 'ti@empresa.com'])

    def test_analista_nao_acessa_solicitacoes(self):
        criar_funcionario(940)
        logar(self.client, 940)
        self.assertEqual(self.client.get('/api/solicitacoes-senha/').status_code, 403)


# ------------------------------------------------------------------
# Regras do Gestor aplicadas no servidor (RF12)
# ------------------------------------------------------------------
class GestorTests(TestCase):

    def test_gestor_nao_cria_nem_altera_gestor_ou_administrador(self):
        criar_funcionario(1000, cargo='Gestor')
        alvo = criar_funcionario(1003, cargo='Administrador')
        logar(self.client, 1000)

        novo_adm = {'nome': 'X', 'cargo': 'Administrador', 'login': 1001, 'email': 'x@empresa.com'}
        self.assertEqual(post_json(self.client, '/api/funcionario/', novo_adm).status_code, 403)

        comum = {'nome': 'Y', 'cargo': 'Funcionário', 'login': 1002, 'email': 'y@empresa.com'}
        self.assertEqual(post_json(self.client, '/api/funcionario/', comum).status_code, 200)

        rebaixar = {'id': alvo.id_funcionario, 'nome': 'Z', 'cargo': 'Funcionário', 'login': 1003, 'email': 'z@empresa.com'}
        self.assertEqual(post_json(self.client, '/api/funcionario/', rebaixar).status_code, 403)


# ------------------------------------------------------------------
# Peças por modelo e fotos da triagem
# ------------------------------------------------------------------
def imagem_teste(formato='PNG', tamanho=(40, 30), nome='foto.png', tipo='image/png'):
    buf = BytesIO()
    Image.new('RGB', tamanho, (200, 30, 30)).save(buf, format=formato)
    return SimpleUploadedFile(nome, buf.getvalue(), content_type=tipo)


class PecasPorModeloTests(TestCase):
    def setUp(self):
        self.func = criar_funcionario(900)
        logar(self.client, 900)
        self.equip = Equipamento.objects.create(pn='PN-9', modelo='Notebook Z')

    def test_salvar_pecas_cria_normaliza_e_substitui(self):
        r = post_json(self.client, '/api/modelo-equipamento/pecas/', {
            'id_equipamento': self.equip.id_equipamento,
            'pecas': [{'part_number': ' l12345-001 ', 'descricao': 'Bateria'},
                      {'part_number': 'L12345-001', 'descricao': 'Bateria 3 células'},
                      {'part_number': 'M99', 'descricao': 'Tela'}, {'part_number': '  '}]})
        self.assertEqual(r.status_code, 200)
        self.assertEqual([p['part_number'] for p in r.json()['dados']], ['L12345-001', 'M99'])
        self.assertEqual(Peca.objects.get(part_number='L12345-001').descricao, 'Bateria 3 células')

        # substituir: some M99 do modelo, mas a peça continua no catálogo
        post_json(self.client, '/api/modelo-equipamento/pecas/', {
            'id_equipamento': self.equip.id_equipamento, 'pecas': [{'part_number': 'L12345-001'}]})
        self.assertEqual(self.equip.pecas.count(), 1)
        self.assertTrue(Peca.objects.filter(part_number='M99').exists())

    def test_mesma_peca_em_dois_modelos(self):
        outro = Equipamento.objects.create(pn='PN-10', modelo='Notebook W')
        for e in (self.equip, outro):
            EquipamentoService.salvar_pecas(e.id_equipamento, [{'part_number': 'BAT-1'}])
        self.assertEqual(Peca.objects.filter(part_number='BAT-1').count(), 1)
        self.assertEqual(Peca.objects.get(part_number='BAT-1').equipamentos.count(), 2)

    def test_listagem_traz_total_de_pecas(self):
        EquipamentoService.salvar_pecas(self.equip.id_equipamento, [{'part_number': 'A'}, {'part_number': 'B'}])
        dados = self.client.get('/api/modelo-equipamento/').json()['dados']
        self.assertEqual(dados[0]['total_pecas'], 2)

    def test_entradas_invalidas(self):
        r = post_json(self.client, '/api/modelo-equipamento/pecas/',
                      {'id_equipamento': self.equip.id_equipamento, 'pecas': 'x'})
        self.assertEqual(r.status_code, 400)
        r = post_json(self.client, '/api/modelo-equipamento/pecas/', {'id_equipamento': 99999, 'pecas': []})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.client.get('/api/modelo-equipamento/pecas/').status_code, 400)

    def test_exige_login(self):
        self.assertEqual(Client().get('/api/modelo-equipamento/pecas/?id_equipamento=1').status_code, 401)

    def test_triagem_so_aceita_peca_do_proprio_modelo(self):
        EquipamentoService.salvar_pecas(self.equip.id_equipamento, [{'part_number': 'OK-1'}])
        peca_ok = Peca.objects.get(part_number='OK-1')
        peca_alheia = Peca.objects.create(part_number='ALHEIA')
        base = TriagemDTO(n_ticket=1, nf=1, cidade='X', estado='SP', id_tecnico=0, id_cliente=0,
                          id_analista_chamado=self.func.id_funcionario, id_analista_triagem=self.func.id_funcionario,
                          id_equipamento=self.equip.id_equipamento, sn='S1', status_peca=StatusPeca.FUNCIONAL)
        prest = Empresa.objects.create(cnpj='33333333000133', razao_social='P', tipo=1)
        cli = Empresa.objects.create(cnpj='44444444000144', razao_social='C', tipo=2)
        tec = Tecnico.objects.create(nome='T', empresa=prest)
        base.id_tecnico, base.id_cliente = tec.id_tecnico, cli.id_empresa

        base.id_peca_trocada = peca_alheia.id_peca
        with self.assertRaises(ValueError):
            TriagemService.salvar(base)
        base.id_peca_trocada = peca_ok.id_peca
        salvo = TriagemService.salvar(base)
        self.assertEqual(Triagem.objects.get(pk=salvo.id_triagem).peca_trocada_id, peca_ok.id_peca)

    def test_dashboard_lista_pecas_mais_trocadas(self):
        EquipamentoService.salvar_pecas(self.equip.id_equipamento, [{'part_number': 'BAT-1', 'descricao': 'Bateria'}])
        peca = Peca.objects.get(part_number='BAT-1')
        criar_triagem(self.func, n_ticket=1, status=1, peca_trocada=peca)   # NFF
        criar_triagem(self.func, n_ticket=2, status=2, peca_trocada=peca)   # defeito confirmado
        criar_triagem(self.func, n_ticket=3, status=2)                      # sem peça do catálogo
        dados = self.client.get('/api/dashboard/').json()
        pecas = dados['pecas']
        self.assertEqual(pecas['itens'][0]['part_number'], 'BAT-1')
        self.assertEqual(pecas['itens'][0]['trocas'], 2)
        self.assertEqual(pecas['itens'][0]['nff'], 1)
        self.assertEqual(pecas['itens'][0]['taxa_nff'], 50.0)
        self.assertEqual(pecas['sem_peca'], 1)


class FotosTriagemTests(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        self.override = override_settings(MEDIA_ROOT=self.media)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.func = criar_funcionario(950)
        logar(self.client, 950)
        self.tri = criar_triagem(self.func, n_ticket=77)

    def enviar(self, arquivo, **extra):
        return self.client.post('/api/triagem/fotos/', {'id_triagem': self.tri.id_triagem, 'arquivo': arquivo, **extra})

    def test_upload_listagem_e_download(self):
        r = self.enviar(imagem_teste(), legenda='Placa danificada')
        self.assertEqual(r.status_code, 200, r.content)
        foto = r.json()['dado']
        self.assertEqual(foto['legenda'], 'Placa danificada')
        arq = self.client.get(foto['url'])
        self.assertEqual(arq.status_code, 200)
        self.assertEqual(arq['Content-Type'], 'image/jpeg')
        self.assertEqual(arq['X-Content-Type-Options'], 'nosniff')
        # aparece na consulta da triagem
        dado = self.client.get(f'/api/triagem/?id={self.tri.id_triagem}').json()['dado']
        self.assertEqual(len(dado['fotos']), 1)

    def test_foto_exige_login(self):
        foto = self.enviar(imagem_teste()).json()['dado']
        self.assertEqual(Client().get(foto['url']).status_code, 302)

    def test_rejeita_arquivo_que_nao_e_imagem(self):
        falso = SimpleUploadedFile('virus.png', b'<script>alert(1)</script>', content_type='image/png')
        r = self.enviar(falso)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(FotoTriagem.objects.count(), 0)

    def test_rejeita_formato_nao_suportado(self):
        r = self.enviar(imagem_teste('GIF', nome='a.gif', tipo='image/gif'))
        self.assertEqual(r.status_code, 400)

    def test_reduz_foto_grande_e_remove_exif(self):
        r = self.enviar(imagem_teste('JPEG', tamanho=(3200, 2400), nome='g.jpg', tipo='image/jpeg'))
        self.assertEqual(r.status_code, 200)
        foto = FotoTriagem.objects.get()
        with Image.open(foto.arquivo.path) as img:
            self.assertLessEqual(max(img.size), FotoService.MAX_LADO)
            self.assertEqual(img.format, 'JPEG')
            self.assertFalse(img.getexif())

    def test_limite_de_fotos_por_triagem(self):
        for _ in range(FotoService.MAX_FOTOS):
            self.assertEqual(self.enviar(imagem_teste()).status_code, 200)
        self.assertEqual(self.enviar(imagem_teste()).status_code, 400)

    def test_triagem_inexistente(self):
        r = self.client.post('/api/triagem/fotos/', {'id_triagem': 99999, 'arquivo': imagem_teste()})
        self.assertEqual(r.status_code, 404)

    def test_excluir_foto_apaga_o_arquivo(self):
        self.enviar(imagem_teste())
        foto = FotoTriagem.objects.get()
        caminho = foto.arquivo.path
        r = self.client.delete(f'/api/triagem/fotos/?id={foto.id_foto}')
        self.assertEqual(r.status_code, 200)
        self.assertFalse(FotoTriagem.objects.exists())
        import os
        self.assertFalse(os.path.exists(caminho))
        self.assertEqual(self.client.delete(f'/api/triagem/fotos/?id={foto.id_foto}').status_code, 404)

    def test_excluir_triagem_apaga_fotos(self):
        self.enviar(imagem_teste())
        caminho = FotoTriagem.objects.get().arquivo.path
        TriagemService.deletar(self.tri.n_ticket)
        import os
        self.assertFalse(FotoTriagem.objects.exists())
        self.assertFalse(os.path.exists(caminho))

    def test_laudo_docx_inclui_as_fotos(self):
        self.enviar(imagem_teste(), legenda='Vista frontal')
        self.enviar(imagem_teste())
        r = self.client.get(f'/api/triagem/exportar-docx/?id={self.tri.id_triagem}')
        self.assertEqual(r.status_code, 200)
        doc = Document(BytesIO(r.content))
        self.assertEqual(len(doc.inline_shapes), 2)
        textos = ' '.join(p.text for p in doc.paragraphs)
        self.assertIn('Registro fotográfico', textos)

    def test_laudo_sem_fotos_continua_funcionando(self):
        r = self.client.get(f'/api/triagem/exportar-docx/?id={self.tri.id_triagem}')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(Document(BytesIO(r.content)).inline_shapes), 0)

    def test_upload_exige_csrf(self):
        c = Client(enforce_csrf_checks=True)
        c.get('/')
        token = c.cookies['csrftoken'].value
        c.post('/api/login/', json.dumps({'login': 950, 'senha': 'senha123'}),
               content_type='application/json', HTTP_X_CSRFTOKEN=token)
        dados = {'id_triagem': self.tri.id_triagem, 'arquivo': imagem_teste()}
        self.assertEqual(c.post('/api/triagem/fotos/', dados).status_code, 403)
        dados = {'id_triagem': self.tri.id_triagem, 'arquivo': imagem_teste()}
        self.assertEqual(c.post('/api/triagem/fotos/', dados, HTTP_X_CSRFTOKEN=token).status_code, 200)


# ------------------------------------------------------------------
# Consulta ao HP PartSurfer (sempre simulada nos testes: sem acesso à internet)
# ------------------------------------------------------------------
class PartSurferTests(TestCase):
    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.addCleanup(cache.clear)
        self.func = criar_funcionario(960)
        logar(self.client, 960)
        self.equip = Equipamento.objects.create(pn='ABC123#AC4', modelo='ProBook 440')
        self.url = f'/api/modelo-equipamento/{self.equip.id_equipamento}/partsurfer/'

    def test_devolve_candidatas_sem_gravar_nada(self):
        EquipamentoService.salvar_pecas(self.equip.id_equipamento, [{'part_number': 'LOCAL-1', 'descricao': 'So local'},
                                                                    {'part_number': 'BAT-1'}])
        hp = [{'part_number': 'bat-1', 'descricao': 'Bateria'}, {'part_number': 'TELA-1', 'descricao': 'Tela'},
              {'part_number': 'TELA-1', 'descricao': 'Tela repetida'}]
        with mock.patch('Daedalus.services.buscar_pecas_por_pn', return_value=hp):
            r = post_json(self.client, self.url, {})
        self.assertEqual(r.status_code, 200, r.content)
        por_pn = {c['part_number']: c for c in r.json()['dados']}
        self.assertEqual(set(por_pn), {'BAT-1', 'TELA-1', 'LOCAL-1'})
        self.assertTrue(por_pn['BAT-1']['selecionada'])
        self.assertFalse(por_pn['TELA-1']['selecionada'])
        self.assertTrue(por_pn['LOCAL-1']['somente_local'])
        self.assertFalse(Peca.objects.filter(part_number='TELA-1').exists())   # nada foi gravado

    def test_sem_pecas_na_hp_vira_404(self):
        with mock.patch('Daedalus.services.buscar_pecas_por_pn', return_value=[]):
            self.assertEqual(post_json(self.client, self.url, {}).status_code, 404)

    def test_falha_de_comunicacao_vira_502(self):
        with mock.patch('Daedalus.services.buscar_pecas_por_pn', side_effect=partsurfer.PartSurferError('fora do ar')):
            r = post_json(self.client, self.url, {})
        self.assertEqual(r.status_code, 502)
        self.assertIn('fora do ar', r.json()['mensagem'])

    def test_limite_por_modelo(self):
        with mock.patch('Daedalus.services.buscar_pecas_por_pn', return_value=[{'part_number': 'X'}]):
            self.assertEqual(post_json(self.client, self.url, {}).status_code, 200)
            self.assertEqual(post_json(self.client, self.url, {}).status_code, 429)

    def test_equipamento_inexistente_e_login(self):
        self.assertEqual(post_json(self.client, '/api/modelo-equipamento/99999/partsurfer/', {}).status_code, 404)
        self.assertEqual(Client().post(self.url).status_code, 401)

    def test_cliente_http_interpreta_a_resposta_da_hp(self):
        def resposta(status=200, texto='', corpo=None):
            m = mock.Mock(status_code=status, text=texto)
            m.json.return_value = corpo
            m.raise_for_status.return_value = None
            return m
        sessao = mock.MagicMock()
        sessao.__enter__.return_value = sessao
        sessao.get.side_effect = [
            resposta(),                                                    # home
            resposta(texto='abc:HPIUID=xyz'),                              # VerifyLogin
            resposta(corpo={'Body': [{'PartNumber': ' L1 ', 'EnhancedDescription': 'Bateria'},
                                     {'PartNumber': 'L1', 'EnhancedDescription': 'dup'},
                                     {'PartNumber': None}, 'lixo']}),     # GetProduct
        ]
        with mock.patch('Daedalus.partsurfer.requests.Session', return_value=sessao):
            pecas = partsurfer.buscar_pecas_por_pn(' abc123#ac4 ')
        self.assertEqual(pecas, [{'part_number': 'L1', 'descricao': 'Bateria'}])
        self.assertIn('HPIUID=xyz', sessao.get.call_args_list[2].kwargs['headers']['Cookie'])

    def test_cliente_http_erros(self):
        with self.assertRaises(partsurfer.PartSurferError):
            partsurfer.buscar_pecas_por_pn('  ')
        with mock.patch('Daedalus.partsurfer.requests.Session', side_effect=partsurfer.requests.exceptions.ConnectionError()):
            with self.assertRaises(partsurfer.PartSurferError):
                partsurfer.buscar_pecas_por_pn('ABC')
